"""
FinePrint — Ingestion Pipeline Tests

Validates:
1. PDF extraction with page map
2. DOCX extraction with paragraph assembly
3. File extension validation
"""

import io
import sys
from pathlib import Path
import pytest
import fitz
from docx import Document

sys.path.insert(0, str(Path(__file__).parent.parent / "backend"))

from ingest.extract import extract_text, _extract_pdf, _extract_docx

def create_sample_pdf() -> bytes:
    doc = fitz.open()
    page1 = doc.new_page()
    page1.insert_text((50, 50), "Offer Letter - Page 1\nWelcome to our company.")
    page2 = doc.new_page()
    page2.insert_text((50, 50), "Offer Letter - Page 2\nTerms of employment.")
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes

def create_sample_docx() -> bytes:
    doc = Document()
    doc.add_paragraph("First Paragraph: Appointment details.")
    doc.add_paragraph("Second Paragraph: Compensation breakdown.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()

@pytest.mark.asyncio
async def test_extract_pdf():
    pdf_bytes = create_sample_pdf()
    full_text, page_map = await extract_text(pdf_bytes, ".pdf")
    
    assert "Page 1" in full_text
    assert "Page 2" in full_text
    assert len(page_map) == 2
    assert page_map[0][0] == 0
    assert page_map[1][0] == 1
    # Check page map offsets match text lengths
    for page_no, start, end in page_map:
        assert start < end
        assert end <= len(full_text)

@pytest.mark.asyncio
async def test_extract_docx():
    docx_bytes = create_sample_docx()
    full_text, page_map = await extract_text(docx_bytes, ".docx")
    
    assert "First Paragraph: Appointment details." in full_text
    assert "Second Paragraph: Compensation breakdown." in full_text
    assert len(page_map) >= 1

@pytest.mark.asyncio
async def test_invalid_extension():
    with pytest.raises(ValueError) as excinfo:
        await extract_text(b"some content", ".xyz")
    assert "Unsupported file type" in str(excinfo.value)
